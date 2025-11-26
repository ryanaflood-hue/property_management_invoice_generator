from docx import Document
import os

def create_clean_template():
    # Load existing template if possible, or create new
    base_path = "invoice_templates/base_invoice_template.docx"
    output_path = "invoice_templates/clean_invoice_template.docx"
    
    if os.path.exists(base_path):
        doc = Document(base_path)
        print(f"Loaded {base_path}")
    else:
        doc = Document()
        print("Created new document")

    # We can't easily "ensure" placeholders without parsing the whole doc.
    # But we can create a text file listing what SHOULD be there.
    
    with open("TEMPLATE_CHECKLIST.txt", "w") as f:
        f.write("Please ensure your base_invoice_template.docx contains the following placeholders in the table:\n\n")
        f.write("{{FEE_LINE_1}}   {{AMOUNT_1}}\n")
        f.write("{{FEE_LINE_2}}   {{AMOUNT_2}}\n")
        f.write("{{FEE_LINE_3}}   {{AMOUNT_3}}\n")
        f.write("{{ADDITIONAL_FEE_LINE}}   {{ADDITIONAL_FEE_AMOUNT}}\n")
        f.write("Total: {{TOTAL_AMOUNT}}\n")
    
    print("Created TEMPLATE_CHECKLIST.txt")

if __name__ == "__main__":
    create_clean_template()
